from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import load_workbook

from health_exporter.db import Database
from health_exporter.exporting import DailyExporter


def test_pseudonym_is_stable_for_private_salt(tmp_path: Path) -> None:
    path = tmp_path / "health.sqlite3"
    private = tmp_path / "private"
    first = Database(path, private)
    digest = first.pseudonymize("account-1")
    first.close()
    second = Database(path, private)
    assert second.pseudonymize("account-1") == digest
    assert second.pseudonymize("account-2") != digest
    second.close()


def test_device_registration_is_pseudonymous_and_updates_metadata(tmp_path: Path) -> None:
    db = Database(tmp_path / "health.sqlite3", tmp_path / "private")
    source_id = db.ensure_source(
        vendor="oppo",
        app_package="com.heytap.health",
        app_version="fixture",
        installation_hash=db.pseudonymize("fixture-installation"),
        account_hash="unknown:fixture",
        timezone="Asia/Shanghai",
    )
    device_hash = db.pseudonymize("android-device:raw-device-id")
    first_id = db.upsert_device(
        source_id=source_id,
        device_hash=device_hash,
        metadata={"android_release": "12", "identity_basis": "android_id"},
    )
    second_id = db.upsert_device(
        source_id=source_id,
        device_hash=device_hash,
        metadata={"android_release": "15", "identity_basis": "android_id"},
    )

    row = db.connection.execute(
        "SELECT device_hash, metadata_json FROM devices WHERE id=?", (first_id,)
    ).fetchone()
    assert second_id == first_id
    assert row["device_hash"] == device_hash
    assert "raw-device-id" not in row["device_hash"]
    assert row["metadata_json"] == (
        '{"android_release":"15","identity_basis":"android_id"}'
    )
    db.close()


def test_daily_csv_and_xlsx_are_rebuilt_from_sqlite(tmp_path: Path) -> None:
    db = Database(tmp_path / "health.sqlite3", tmp_path / "private")
    installation = db.pseudonymize("fixture")
    source_id = db.ensure_source(
        vendor="oppo",
        app_package="com.heytap.health",
        app_version="fixture",
        installation_hash=installation,
        account_hash="unknown:fixture",
        timezone="Asia/Shanghai",
    )
    run_id = db.start_run(source_id, "test", "fixture", "fixture", {})
    db.upsert_daily_metric(
        source_id=source_id,
        canonical_metric="heart_rate",
        local_date="2024-01-02",
        summary={"min_bpm": 51, "max_bpm": 143},
        chart={"series": [{"timestamp_utc_ms": 1, "value": 51}]},
        fidelity="exact_source",
        is_partial=False,
        run_id=run_id,
    )
    db.connection.commit()

    manifest = DailyExporter(db, tmp_path).export(
        {"csv", "xlsx"}, source_id=source_id, output_dir=tmp_path / "export"
    )
    assert manifest["tables"]["daily_overview"] == 1
    with (tmp_path / "export" / "heart_rate_daily.csv").open(
        encoding="utf-8-sig", newline=""
    ) as handle:
        row = next(csv.DictReader(handle))
    assert row["min_bpm"] == "51"
    with (tmp_path / "export" / "blood_pressure_daily.csv").open(
        encoding="utf-8-sig", newline=""
    ) as handle:
        empty_table_rows = list(csv.reader(handle))
    assert empty_table_rows == [
        [
            "source_id",
            "date",
            "revision",
            "fidelity",
            "partial",
            "min_systolic_mmhg",
            "max_systolic_mmhg",
            "average_systolic_mmhg",
            "min_diastolic_mmhg",
            "max_diastolic_mmhg",
            "average_diastolic_mmhg",
            "normal_count",
            "high_count",
            "low_count",
        ]
    ]
    workbook = load_workbook(tmp_path / "export" / "health_daily.xlsx", read_only=True)
    assert "daily_overview" in workbook.sheetnames
    assert "heart_rate_daily" in workbook.sheetnames
    assert list(
        next(workbook["blood_pressure_daily"].iter_rows(values_only=True))
    ) == empty_table_rows[0]
    workbook.close()
    db.close()


def test_partial_day_updates_create_versions_without_deleting_history(
    tmp_path: Path,
) -> None:
    db = Database(tmp_path / "health.sqlite3", tmp_path / "private")
    source_id = db.ensure_source(
        vendor="oppo",
        app_package="com.heytap.health",
        app_version="fixture",
        installation_hash=db.pseudonymize("fixture-installation"),
        account_hash="unknown:fixture",
        timezone="Asia/Shanghai",
    )
    first_run = db.start_run(source_id, "snapshot", "fixture", "fixture", {})
    first_id, first_created = db.upsert_daily_metric(
        source_id=source_id,
        canonical_metric="activity",
        local_date="2026-07-22",
        summary={"steps": 1200},
        chart={"series": {"steps": []}},
        fidelity="exact_source",
        is_partial=True,
        run_id=first_run,
    )
    second_run = db.start_run(source_id, "snapshot", "fixture", "fixture", {})
    second_id, second_created = db.upsert_daily_metric(
        source_id=source_id,
        canonical_metric="activity",
        local_date="2026-07-22",
        summary={"steps": 2400},
        chart={"series": {"steps": []}},
        fidelity="exact_source",
        is_partial=True,
        run_id=second_run,
    )

    rows = db.connection.execute(
        """
        SELECT id, revision, summary_json, is_partial
        FROM daily_metrics
        WHERE source_id=? AND canonical_metric='activity' AND local_date='2026-07-22'
        ORDER BY revision
        """,
        (source_id,),
    ).fetchall()
    current = db.connection.execute(
        """
        SELECT id, revision, summary_json, is_partial
        FROM daily_metrics_current
        WHERE source_id=? AND canonical_metric='activity' AND local_date='2026-07-22'
        """,
        (source_id,),
    ).fetchone()
    assert first_created and second_created
    assert [row["id"] for row in rows] == [first_id, second_id]
    assert [row["revision"] for row in rows] == [1, 2]
    assert [json.loads(row["summary_json"])["steps"] for row in rows] == [1200, 2400]
    assert [row["is_partial"] for row in rows] == [1, 1]
    assert current["id"] == second_id
    db.close()


def test_coverage_export_honors_source_filter(tmp_path: Path) -> None:
    db = Database(tmp_path / "health.sqlite3", tmp_path / "private")
    source_ids = []
    for name in ("first", "second"):
        source_id = db.ensure_source(
            vendor="oppo",
            app_package=f"fixture.{name}",
            app_version="fixture",
            installation_hash=db.pseudonymize(name),
            account_hash=f"unknown:{name}",
            timezone="Asia/Shanghai",
        )
        run_id = db.start_run(source_id, "test", "fixture", "offline", {})
        db.set_coverage(
            source_id,
            "heart_rate",
            "2026-07-01",
            "no_data",
            {"database": False, "ui": False, "no_data": True},
            run_id,
        )
        db.finish_run(run_id, "complete")
        source_ids.append(source_id)

    tables = DailyExporter(db, tmp_path).build_tables(source_id=source_ids[0])

    assert len(tables["coverage"]) == 1
    assert tables["coverage"][0]["source_id"] == source_ids[0]
    db.close()
